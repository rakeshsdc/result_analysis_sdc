"""
FYUGP / University Mark-list Result Analysis App
--------------------------------------------------
Upload a University "Mark cum Grade Statement" PDF (one page per student,
in the standard University of Kerala FYUGP layout) and generate a full
Result Analysis Report — viewable on screen and downloadable as a
formatted Word (.docx) document with a college-logo letterhead.

Run locally:
    streamlit run app.py

Deploy:
    Push this folder to a GitHub repo (app.py + requirements.txt) and
    deploy on https://share.streamlit.io (Streamlit Community Cloud),
    pointing it at app.py.
"""

import io
import re
from collections import Counter

import pandas as pd
import pdfplumber
import matplotlib.pyplot as plt
import streamlit as st

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# --------------------------------------------------------------------------
# The official grade hierarchy used by the university (best -> worst).
# "F" here represents an overall Fail result (the mark sheet prints "--"
# for Grade Awarded when a student fails, so we relabel it "F" for the
# purposes of grouping/ordering in this report).
# --------------------------------------------------------------------------
GRADE_ORDER = ["O", "A+", "A", "B+", "B", "C", "P", "F"]

# --------------------------------------------------------------------------
# Fixed institution identity (no longer user-editable).
# The logo is embedded directly as base64 so the report always has it,
# regardless of how/where this app.py file is deployed (no separate
# "assets" folder or file path to keep track of).
# --------------------------------------------------------------------------
import base64

COLLEGE_NAME = "Sanatana Dharma College, Alappuzha"

COPYRIGHT_NOTICE = (
    "Developed by Dr. Rakesh Chandran S. B., Associate Professor, "
    "Department of Physics, Sanatana Dharma College, Alappuzha. "
    "An IQAC SD College initiative."
)

_LOGO_BASE64 = (
    "/9j/4AAQSkZJRgABAQEAYABgAAD/2wBDAAMCAgMCAgMDAwMEAwMEBQgFBQQEBQoHBwYIDAoMDAsKCwsNDhIQDQ4RDgsLEBYQ"
    "ERMUFRUVDA8XGBYUGBIUFRT/2wBDAQMEBAUEBQkFBQkUDQsNFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQU"
    "FBQUFBQUFBQUFBQUFBT/wAARCADHAO4DASIAAhEBAxEB/8QAHwAAAQUBAQEBAQEAAAAAAAAAAAECAwQFBgcICQoL/8QAtRAA"
    "AgEDAwIEAwUFBAQAAAF9AQIDAAQRBRIhMUEGE1FhByJxFDKBkaEII0KxwRVS0fAkM2JyggkKFhcYGRolJicoKSo0NTY3ODk6"
    "Q0RFRkdISUpTVFVWV1hZWmNkZWZnaGlqc3R1dnd4eXqDhIWGh4iJipKTlJWWl5iZmqKjpKWmp6ipqrKztLW2t7i5usLDxMXG"
    "x8jJytLT1NXW19jZ2uHi4+Tl5ufo6erx8vP09fb3+Pn6/8QAHwEAAwEBAQEBAQEBAQAAAAAAAAECAwQFBgcICQoL/8QAtREA"
    "AgECBAQDBAcFBAQAAQJ3AAECAxEEBSExBhJBUQdhcRMiMoEIFEKRobHBCSMzUvAVYnLRChYkNOEl8RcYGRomJygpKjU2Nzg5"
    "OkNERUZHSElKU1RVVldYWVpjZGVmZ2hpanN0dXZ3eHl6goOEhYaHiImKkpOUlZaXmJmaoqOkpaanqKmqsrO0tba3uLm6wsPE"
    "xcbHyMnK0tPU1dbX2Nna4uPk5ebn6Onq8vP09fb3+Pn6/9oADAMBAAIRAxEAPwD9U6KKKACiiigAooooAKKKq317BYWstzcz"
    "R29vGu6SaVtqqvruoAtUV4J4y/bk+BXgO5FpqHxJ0e8v92xbPRWfVJmb+7stlfDexrl0/bU1vxb8nw++AfxM8Ts3MV3q2nxa"
    "LYy/7s87/wDstAH1HRXy4njv9rHxTCJdP+Gfw78AR/xJ4q8RT6i6r6/6Gm3/AMerN1Lwv+0deSSp4h/aC8EeBJIo0mkg0Pwx"
    "FceWrPsVv9Ml+6zjarH+Lj2oA+taK+TNJ+A3jfxZrusaJqP7VvizU9Y0gwrqVlodnp2ny2vmrvi3qkTbN6/MtcNb/DXwDqn9"
    "pyT/ALUfxvmexLefv1+4skZVuEtXeLbap5qJOyxM8W5VZvmoA+7KK+Htd+Bnwx8OwJLq/wC0J8YoR/b6eFvMl8XXn/IRYbhD"
    "/qv7rbt/3F/vVBefBv4a6P4qv/DMnx++Oun6nY3UlrcS/wDCR6ilp56WX2xovtX2fyGkW3zLs37qAPuiivhfwv8ADvwT4iVp"
    "vDP7V/xelVdDXxQxvdf85YtO3MPPfz7X5F+X7rfNx92uth+FPxI0vwbH4r0X9rvULTwrNbJexat4g0HS7y0+zsuVdpX2fL8y"
    "/wAS0AfXdFfJel2/7Tttp9nf+HPi18KPiLp16M2V5rOkz2SXX+41nK6t/wAB3Vpr8XP2nPCrbNe+BGh+Loh9+88I+LYoAP8A"
    "dhukVm/76oA+oaK+XP8Ahvbw94bk2fEP4c/Eb4aIv+svta8OS3Fjn/ZntvN3f9816b8Ov2ovhH8WVhj8KfEXw/rF3N9yxS+S"
    "K7P/AGwfbL/47QB6tRRRQAUUUUAFFFFABRRRQAUUUUAFFFFABXyb+2x8RviPpfir4V/Df4eeILfwVeeO7y9t38UTw+f5EkES"
    "SxQKv8JlZtu6vrKvlH/go54evV+ANt480dT/AG78O9esPFVps+8fKl2Sr/u7JWb/AIBQB8yfFz/goh42tvCvgfwvL4kPwx+I"
    "Gm6q2jfESKPQvtt1ZorIn2+13r5Bif522ff+dNvy/M3p/wAVPgH8MPh9/wAIxqnxPuPHnx0TVma4XVvEWvSvplmi+V0WJkiT"
    "csm5Eb5W2P8AN8tcJ+0p8QPAnxO+M3h/xN8GdP1L4peLdRsl0bxv4b8PWMtxZanok8X3bq4QeXFOnyhW/hZU37fKWtPwj8Sr"
    "L4E+HdV+A3xe8H+NvHNz4Pm+0eE7jwzZzy3V/o8kb+Ud8Drt8pN8T/Ps2/J82ygD0KX4zeHfhZr3xJ8GfC7wb4V8MS2OnL/w"
    "i2t6Ppe621C+awgv0t52iVU/frP+6+f59j/7G67qHjn49/FLQW03R9G1Lw34lXxNa3dnqX+o0w6c2lPOn+keVmW2a8iXcjL5"
    "uyZYn2vVb4RfHT4sfGzwRHqvwL+EHgrwP4OupdkGs+KNWVkl8jbb/NZWS70dViRBvb7sSfw7K7k/s9/Hrxyrf8Jr+0TeaPay"
    "/wCs0zwJoNvYeX/uXUm+WgDG1z9n7xd410vx1Kug2fhnUPGmi6TaXf8AbF8dQSxl+2y/2vbrsl3NbtbrEyKrru+7+6/g84+I"
    "Pwl8OLptpD4z/aB8H+HfE2j6E3h7Ttf/ALQiinnt0uElgS4gll+Zdm5WTezbkil37q9kX/gnn8K9WVH8ZXfjD4lXCHeJvF/i"
    "i8uD/wB8o6L/AOO13fhz9j74H+Eo1/sz4T+EUdfuyXOkxXEo/wCByqzfrQB82+H/AI3fs+/C34xap4/m/aCtdS1fVru+k1Ky"
    "sbc3FrPbyP8AuIl+zxM+6FEt1V3d/uPtVfN+Xk3+LH7OXiWw1Ozu/iJ4o8WaPMuqW1jY2PhG88qzivtQiv7j5vsv7397FEvz"
    "/wAK/d+av0H0nwjoXh1VXStE07TkX7q2dokQ/wDHVrboA/N/W/En7NuvQ3unDRfihbaDLcX91Houh+HL/TbS1lu4reKV0SCK"
    "L7qQfKrbl/ey7lbdVi++IXwCbxFqXiOd/iVaeKb5bqO/8QXHhGf7RPFPYRWbK3+i/c2RLKv8Svu/gdkr9GKKAPzE0vxd+yr4"
    "Z/t6x/4W14m0K212Zv7QttQ8P38TeR56Tvao7Wu2KL91Emz+5uX+OvUtF+LnwPvPhJoXhDQvj34Xmfw/rq65pkviV/IhZUu3"
    "nhtbhXZNyoHCqybdnlo235Np+58blw1cxr3wy8H+KE2614T0PVg3a+02Gf8A9CWgD5W8QfC22+NENzD4a1/4XeM59X8PNo7X"
    "ljdLEmgXUtxPcT6hp8ESz75X+0I3+tibfbxNv+b5ew1C3+Jvh3wX4zsIfDviL+1tQ8XrfnWrC7tZS2lvqESv5CrL5vmrYQ7d"
    "uxfmaun8UfsJ/AHxgG+3fCfw3Bu/i0y1+wN/5A2VzB/YL8NaCwbwH8RPiZ8OwvMVtonimWW09t0Vx5u5fagBdc+LfjLVP2hU"
    "sbG51TQvBVvplnOun3Vj9nuNRm2S3NxsWeybzP3XlRbEuInR0f5X/h8W8UeOPhl8UvDupa18QPhr4F8Z3F1pml38FnollPFq"
    "FvJqMuyG3a82fvfl812liddv2eX91/FXszfDH9qHwHk+GfjJ4X+IMS/csfHXh42jIvp9os23Of8AaZa8i+If7QD/AA48S6L4"
    "N+NP7PLW+p61ePqWnaj8MdTW7uL26hVle6iii8q5iZElb592/wCZ6AKnjr4d6d+z7+zzqnxc+HvxQ8cfCmztbZm07w3LrEHi"
    "DSrmXcUgjiilaVW819vzbtyp95V2tWH4d/a++J/7Rvw58I+B/A3i2x0jx7Fp8ut+PvGkOlSxWnh22iZmW38qVPnnb5FfZ8vy"
    "tt+XcyQ/E34ieJv2nNdk8U+BPBXiPxH8PvhHawTaRoep20puda8ROqrG86SnzZVtVbcwyzMyfxebXI2vxc8F/Dr9if4heGtF"
    "1nULz4/+MpY7LxPYeILSWz1i71G/fy5R5Uo3NGkUsqrt/wB75WegD7t/Y3+LmvfHT9m7wV438TWcVrrmqQS/avJXakpinli8"
    "1V/h3+Vu2/7Ve31x/wAJ/Att8L/hn4U8H2RV7fQdMt9PWRf4/KjVN3/AsbvxrsKACiiigAooooAKKKKACiiigArnfHvg+y+I"
    "XgfX/C+ojOna3p8+nXW3r5UsTI36NXRUUAfDfwK0n43/ALHngO2+HVn8E9N+Img2M8zQeJfCuu29hNeb2Z91xbz/ADeZ/Du3"
    "fwotd1e+OP2pPidDJpmh/DXQfg3DL8reIPEuuRaxcIndorW2XaJP+urba+qqKAPKv2d/gfo/7Ovww03wZpN1daoYJZbm91O8"
    "OJr26lfdLK393PGF/uqv3vvV6rRRQAUV518WPj94A+BumrfeOPFWn6CknMFvK++6uR0/dQJukl5/urXk8f7QPxg+KjlfhV8H"
    "Z9J0xuI/E3xMnbS7fPUMtlHvuJVbs3yUAfTtFfLepfBP4ta9Zy33xL/aMvPD2mYDyWPgjTrXRobf/t6nEsrD/e21xH/Cj/2Z"
    "de1D7Bq/jPxL8YNZSL7Q0LeKtU1+4Me/7zRWbt8u7/ZoA+zbjVbKzOJ7u3hP/TWVVqa1vbe8XdBNHMP+mbhq+RLP9mf9l21j"
    "2W/wT1KYf9NvCGsn/wBGxVy198I/2O7rW9Y0d/htqWnaxo9st3fwWvhrXraS0gfdslZoohtT5H+f7vytQB92UV8SaL8JfgJe"
    "WNjceBfjf4u+HMV6iXFjHp3jm6t0uFb7rLDfO/mr/wABrvk+Hn7RngmNbjwj8XfD/wAStNx+603xxoy2zlP9m9s+Xb/eioA+"
    "naK+Xv8AhsbUfhuwh+N/wy8RfDaEMqt4isf+JzoWP7zXUC7otx/hdK+hPB/jTQfH2h2+teG9Zsde0q4/1V7p1ws8T/Rl4oA3"
    "a+ff2hP2c9Y+I3i7wp8RvAXiSHwr8TvCqyx6fdX1ubixvLeX79rdJ97YefnT5l3N/slfoKigD5dj+OH7R3hz/RfEH7OcPiCa"
    "L5TqnhXxbai3nb1WKfbKi/71cLqHwR+JH7Unx8+HPxA+I/w40b4Y+H/A119vt7U6nFqWr6nKArxI8sX7tYklRG2N/t/3q+3K"
    "KACiiigAooooAKKKKACiiigAooooAKKKKACiiuS+I3xE0H4U+DNW8V+KNRg0nRdLgNxcXEzdB/dXn5mY4VV6s2KANfXte03w"
    "zo15rGsahb6VpdnE01ze3cqxQwqvVnZvu18w/wDC6PiX+1NJNZfBKAeCfAG9opvibr9nulvB/F/Zdm3+t/66y7V+9/EtZfg/"
    "4c+KP20NVsvHPxXsbjQfhXDKt14b+G9wBvvlz8l7qn97d95IPu9P+B/X1raw2NvHBBEkMMa7Y44kCqq4+6tAHjvwj/ZQ8A/C"
    "PUJNdjtbnxR43nO658YeJpjf6rM+Of3r/wCr/wB2LaOK9fvrVb+0nt2kkhWWNoy8LbXXdxlW7NV2vNvi58fvAfwN0+O58Y+J"
    "LbS5bg/6Jp67pby7bptigQNI/P8AdWgD50/aM+FOnr4p8M6efBmtWMEV3HfwfFCzbUvEOp2LxPEz2/lRLLP+9TdErSt5W3f8"
    "v3VrvfD3w68St+0t4Z+Ktne6xrmhax4c1DQ9TbWLGDTZdMiW6insgtu0UU53N9o++rsN38O6q0Xxm+O/xah3fDj4T23gbR5O"
    "Ytd+KF09vKy5/h0633Sr/s73WrUf7O/xi8V5fxn+0VrtvG3LWPgrRrPSkiPosrpLK3/AqAKfgX9knTb3wHq2gfE/QdD1zxLM"
    "ksK/EC3lafWLxpGfbdebLFvtZ0+QqqO6o33ePlrO8Xfs3/ELw/8AC3x5YeCtT8I6h4j1vQpdKa/uNBuF1K+XymRInupb9lX7"
    "7/wbFZvuV0K/sT6dNGv274vfGLUnX+K58b3Cf+OxBFpD+xeLcb9L+OHxk0tl+6v/AAlrXSL/AMBniegCHx/4S+LHgPw34T8O"
    "/Dy1fUPDOjeGP7NTT7CKwe4lv4kRLf7R9sdU+zbE+byn37v71eDeH/gv4z+C8fwY+EGj3thpXi7XpW1TWvE3hma4tb61s4GS"
    "e/adNzwT7nl8hJZfl+ZNsX9z3uT4R/tF+B1M3hP426Z4zRB+70jx94eiVG/3rqz2Sf8AjjV5h4w+Iur+C/iBZeOvi38OfEvw"
    "z8UWVmunRfEPwjdy69oD2au8rxXkCfNFAzv/ABxb/wC66bFagD7L0TRptN0WGwv9TuNdlVWD3V9FEskyn+8sSIn3fl4WvBfG"
    "X7HOlWGuXHi74O61c/B/xs53Svose7StQYEfJeWH+qdf9pQrclvmr3Hwb488PfELQbbXfDGtWGv6RPxFfafcLNE2P4dy/wAX"
    "+zXSUAfNHgX9qLVvDPiyw8AfHXRLfwD4yum8rTNes5Gfw/rrf9Os7/6qT/plL83T+8Fr6Xrj/iN8NfDPxa8H3vhjxdo1vr2i"
    "XiFZrS6TK57Ov8SOv8LKdymvmzQ/G/if9ivxTpng74i6pd+J/gxqlwtn4e8cXx3XGiSMfkstRb+KL+FJ+38Xy/cAPsOiokkW"
    "RFZW3K33WWpaACiiigAooooAKKKKACiiigAooooAKKKKAInkWNGZm2qv3mavjXwrZv8AtzfGBvFmrL53wH8E6g0Xh3T5F/de"
    "JdUi+WW9lX+O3ib5UX7rt/wNa7D9tzxdq+q6P4V+DHhK8az8WfE68fTWuo/vWOmRrvv5/wDv18v/AANq968C+BtG+HHg/R/C"
    "/h+1Fjo2lWyWlrbqSdqKMf8AfX+1QB0tFFfCP7cf7WWg2XjrTfgXbeO4vAcWpL5ni3xQoZ5dOsmXP2WDYjfv5VP/AAFWX+98"
    "oB6J40/aG8YfGrxnqXw7/Z/+ytNp0pg8QfEbUIvO03Rm/wCeVun3bq6/2PuLld38Wzvfgv8Asr+Dfg1qEuurHd+LfHV58194"
    "x8SSfatSuG4DbXb/AFSf7Kbf+BV5j8N/2xv2VvhH4L0zwp4V8cabpOiadEIYLeHTrznvvb918zNyzN/FXV/8PEv2dv8Aoptl"
    "/wCAN5/8ZoA+kKK+b/8Ah4l+zt/0U2y/8Abz/wCM0f8ADxL9nb/optl/4A3n/wAZoA+kKK+b/wDh4l+zt/0U2y/8Abz/AOM0"
    "f8PEv2dv+im2X/gDef8AxmgD6QpuNy4avnH/AIeJfs7f9FNsv/AG8/8AjNH/AA8S/Z2/6KbZf+AN5/8AGaAK/j79kf8AsnxH"
    "ceOfgprCfCzx+3z3MNtEP7E1j/YvLNfl7/61BvXcW+Zq3fgT+0sPH/iC9+H/AI30Z/A3xZ0iMPe+H7ht8d5F/wA/dlL/AMt4"
    "G/76T+L+8cr/AIeJfs7f9FNsv/AG8/8AjNeO/tHftFfs1fHLw3bT2fxdtfC/jzQ5TeeHfFVnYXXn6fcqf+uPzRPgKyfxLQB9"
    "5VgeMvBujfEDwtqnhzxBYRanoupQtbXdpMp2yI3b/wCv7V4h+xR+1Xp/7Unwya8lmtl8X6JILLXbW1Y+UZf4biL/AKZS7Sy/"
    "8DX+Gvo6gD5K/Z78U65+z98UH/Z68b6lLqNg8LXvw/8AEN4ctf6en37GVv8AnvAP++l/u/KD9a14N+118Fr34yfCiWfw25s/"
    "H/hi4XXvDF/CfnjvoPmVP92VRs/4Ep/hrqf2dPjLZfH34N+GfHNkiQNqVt/pdrzm2uU+SeLk/wAMisP93bQB6hRRRQAUUUUA"
    "FFFFABRRRQAUUUUAFFFVri6js7aWeVtkcSM7N6KvWgD5V+CMP/C3/wBs74yfEacGXTfBiweAtEf+ESJ+/wBQ/wCBCVlX/dav"
    "rOvlT/gmnbSXP7LOneJ7pcaj4v1nVNfvG/vSy3cqf+gxJX1XQB578ePitYfAz4P+LfHmoqskGiWL3CxM23z5fuxRf8DlZF/4"
    "FXy74T/YqvvEP7MNrqOpXraZ8eNS1AeO/wDhKGTZcWutv+8SJv7sSrtiZPu/eau0/bojXxjqnwO+GMm57XxZ45tZtRhAwJ7G"
    "zVriaI/+Q/8AvmvrCgDxf9mL49D49fDv7ZqFk2heNNFun0jxLoMv+ssb+L5ZVK9kb7y//YtXtFfIX7SGg6n+zl8UrX9ovwna"
    "zXWjNHHpvxB0e1Xe15pufkv1XvLb/wDoH91d9QfF74seNYdfuPEOj+Ir3TfCV3bwXXhTXLdrVPDzRvaq6PfyvE8srTXR8jyk"
    "2vseJokZ97IAfX0is0beWwV9vyuw3V8n6l8ffi5peialIi+G9Z1SPx0ng+2t7DQJo3ZPK3vcbJtRRWZv4V81FH95q80/ab1q"
    "7n+IHhn4paR8XPHOm/Be/u10LxPD4X1ZrWTw9ebdsVw0UsTeVEdy+ajqrLuVv49tek+PP2e9J+GvgXVfF+v/ALRnxqtfD+l2"
    "32y6u08RxSjy/wDZVLX5v+A0AZTftmeI9K+H9xqurW3h+w1qXwlrGpadaXSS27XWs2t08EVh5LTbvNYGLfaqzOrttV3+9XY/"
    "Ej4zfEvwFqXxXlWTw3d6d4R8HReJ7S0Oi3K3M8sxv1SB5ftm3EX2JdzIvzb/AOCvzh+Nn7amgaTb3Nn8LPi18c9e1JW2x6lr"
    "3iCC3svvfeWNYPNf5f72yu9/ZK/aA0L436tpfhDxx8bfjN4S8bXrPFFcQ+KIBpkxVHbO+SDdG2F27W3ZY9aAP0m+E/xS1vx9"
    "408eaTq/hufw1DoE9nHa2l60T3bLLb+azStFLLE3zH5djdua9Xr49+K3wH0v4N/DfXPG/iL9on40W2i6TavPK6+KIN8v9xE/"
    "0X5mZtqL/vVc/Y1tfFnw98C6NdfF7xd4kv8AxR8QLxm0nSvENzLdf2dEkTyw2rPt2pO0W933bN23bt+SgD62r5W/aw8e658Q"
    "PE2k/s9fD2/Nn4p8VQG48Q6xB10LQ87ZZf8ArrL/AKpP97+HejV6x+0B8cNL/Z/+GOo+KtRhkvbvzFstK0uDJl1G+l+WC3jX"
    "1ZvToqs3auR/ZP8Agfq/w18Pax4r8cTpqPxS8aXA1TxHfr0jbH7qyiHaKBDsXqPvY+XbgA8i+Mvw60D9i7xp8KPip4Lsl0Xw"
    "dpfkeCfFlrD92TTJ2/0e7k4+9FOdzP8AebetfbleW/tNeAYfid+z38RfDFwgmOoaHdLAD/DOsZeF/wDgMqI34Vm/shePpPiZ"
    "+zD8MvEdxMZrq60O3iuZmP354l8qVv8AvuN6APZK+Sv2b0Hwk/aq+NvwnRDDo+rNb+PdDi/hWO5/dXu3+6ouFUKK+ta+UvjV"
    "IPCP7ev7POtxfIPEmla94evH9UiiS6iX/v7QB9W0UUUAFFFFABRRRQAUUUUAFFFFABXP+PldvA/iJY/9Y2nXAX6+U1dBUMka"
    "zRtG67lZdrLQB86f8E7Gik/Yt+FZjxt/s6Rf+BfaJd365r6Rr4//AGAtbuvBv7LOt+GP7Ou9a1X4d69rehSababftFxLBcNO"
    "sSb2Vd7ecqrvZRW/8QPix+0NqXgvWrjwp8GIvC17bQG7trnWvEFtdTv5Z3+UtnapL5juq7dnmp9/7y0AU/2kWCftofsmyyc2"
    "4u/EyMzfd3NpqbK+ra/Pj9qzWvFnwz8ReBfEHi3WofFf/CvPEWj+Jpdbh05LCZdMvJZrO9ieJX2/K6W+1/8Aprtf7m5/0CSR"
    "ZEVlbcrfdZaAK1/p9tqVnPaXkEdzbTxtHJBMu9ZFbhlZT94c18OeGfBC/Av4sWXwG8Q6vqGmeCtS1M+JPhZ4iidWbTLxd3m6"
    "W3mqyPt819ivu3K/9512/eNeVftFfA3Tvj98M73w7dTnTtUikS/0bV48iTTb+I7oLhSP7rf+OswoA+bvAni2OTVvEPhLxZe6"
    "L4tt/GmuX9p4r8P6o7trtpaxRPb/AG+9eLbBbQeVaxbU8qJdjptld/veZ+IvhP4/8ZabJ+yNffEhfDdlbKNT8M6rfWaXKeJP"
    "D4+5as27Pn2rKvyL95U3fdVXr0LQNX134/eF5LnxNa2tx8Tvhxu0PxR4F1Wwn1XT7q63q1vqUVlFLEsjNhmRn+X52+ZNm+sr"
    "wzpMX7Rnwlh8JeHLk+Hfif8AC+5eXwZqr6hDcXjSW2xJXuEg3xW9tK58jy/NdWWL5dyoNwB84/tOf8EqdM/Z/wD2f9f8e2/x"
    "Buda1LRFikmtZNMWCKdWmSL5f3rMn3938XSvQtV/4Ihwt4ftG074qSJraopuVutI3W7t/Fs2y7l/8er1v9oz462/x7/4Ju/E"
    "zWJrJtH8T6clvpfiDQ5D+907UIr6BZYm/wBnPzL/ALLV6r+1l8Stf17WtE+BXw6vPI8eeM4nfUNShY/8SLR1Oy4vHx9xm+aK"
    "P/a7httAHhvwZ8F3XxdvNC0XXvEN943+C3wMHlLqMOmM/wDwlOsW6fLst4t7TQ2qYRFXe0rf399eq+A/Dun/ABovrjxJ8QNA"
    "8M+O9MvriS6j8XaFqqr/AMI15Sb1spVfypYHt/u+an73e250i/hxviBdp8K/BXw8+H3w88I6drfwzZ7fS9H13T/Edxpstvq3"
    "2poN0t1AjeVuZ5X835lldZU+86K/N+MvCupfEL4gT/ATQtZ1C517xFHbav8AFnxM93FPLZ6ckSJDpqSxW9uu+VPl/wBUrbH3"
    "tv3vQB2PwP0WX9qn4yW3xf1SO4f4Y+C3k0z4e2d47y/2jKuY7jV5d/zO3y7Imbn5d3Drub7KrH8OeGtM8H+H9N0PRbKHTtI0"
    "2BLW0tLddqQxKu1UWtigCnqUkcOn3UsrKsKxszM393bzXyF+wN420z4a/sA/DTWfENz9j08TvbNcMNqRm51iWCJnZuAu6VNz"
    "fwr9K9m/bA+IkXwq/Zl+I/iOSTy5odGnt7Uj/n5nXyIP/IsiV4vqXw/1fwH+zj8DPhRYeEbHxjrNvFa32r+G9Su0tbaeK0i8"
    "268x3jZdq3Utvwy/NuFAH1Ovjrw6+k3WpprmmzadawtcT3kN3G8UUaruZ2bd933r5c+PmvweLPjh+yVq1nBc2seo6/f3ltHc"
    "KFlMBsy+5lPK7k2Ns+983zfNxXF/st/Bz4R/G/Q7yy8ZfADTdB168jbxZBcSwWrRT6dqN3dS2XlSwPuVEiQR7GVPufdr0b4i"
    "afb6/wDt8fAjwtYxhbbwN4W1fX5Yuu2GdVsIf/HloA+taKKKACiiigAooooAKKKKACiiigAooooA+Sfhyx+C/wC3h8QfCUye"
    "ToPxO02LxXpHHyfb4F8q9iX/AG3Uea3/AAGuU8ffDf4hePfFnjXWPEHhSbxx4UvPEscb+DtNMVrPFBpxdLR3M8sSXEV1FK8r"
    "tv8Akb7PtVtrLXqf7anwx1zxZ4B0nxr4It/M+I3w9v18QaGqKd9yq/8AH1af7ssX8P8AEyItX7XSvCn7anw28FeMbbxP4ks/"
    "Cd5A81zoWias1hHes2Fkt7xosSt5TI67VdOd1AHn/wAIPhXafEvwz8QfBnjvxAmq+KZPC2neE9Y0dZPtEukwbbmWL/SC7efL"
    "/pBVpfl/e2jf71dN+xF8SdU1XwTqHwx8ZS+X8Q/hvMuhapE7/NcwKP8ARLxO7JLEF+Y/3f8Aary/4VeJl0LTfDnjDwF8E4/A"
    "GhLPf29j/ZM8Uv8AbNnEzfaLe/VUVoLr/RWliaVmTfF5Tyr5vzdx8ZPBd/4vk8L/ALRPwJmttW8Y2Fhl7GKXFv4p0hjuazc/"
    "89VPKE/db5ey7QD6xorzf4G/HPwx8f8AwTF4i8NTsAj/AGe+026Xbd6ddL9+3nj/AIHU16RQB8pftVeA9a+GvizSv2h/AlnL"
    "fa94btza+KdFtVwdc0TO6VP+usP+sQ/7Hfaq15ZrXhvwzoum2vjvSPFd7p/gCSzuvEPgvTvCdssUSRRWD3TIlrEEiW6gktX/"
    "AHtxuWVbhon+avv7G5cNXwbrXw50X9n34lXPwt17S9Nuvgv8RruWfwhNq1ol1beG/EDq2612P92KXduRf95f79AHln7bngzX"
    "fF3wt8YfFz4Y6XdW1p4oVPDvj3wz5Yl8+S2uk+z6hH5TMkjI6Km9d3yS/wC/Xp0OneMv2c/g/rfxE8X6G+v/ABV+IvmXXijV"
    "hBcPb6JarbP9n0/fA/mxLxHAsm9VRnZ3l/dJu9x+BXwn+Iug+MtT8TeMfEDW4mAtT4etr5r20kiW1sokdfkiigCy293IqRRL"
    "8t38392vc9e1/T/C+h3+savdw6dpVhC9zc3Vw21IolXc7t9KAPz+8O+Kl+CXw1HjPR7i+1jUr/VW0HwR4LS7jv01bUUiSztZ"
    "Yr1Fia5sYIlZER4lVW819zs0UtfVP7LPwHPwL+H80eq3v9t+OfEFw2r+J9ccfPd30nzP83/PNPur/wB9Y+avLP2ZvDWoftCf"
    "EqX9oXxLYyafoEcMum/DrQrhdn2PTc7Xv2X+CW4/9A/vLsNfYFABRRXzx+0l+0Ze+A7qx+Hvw6s4vEvxi8QKV03Ss7otNjPD"
    "X15j/VwJ975vvYoA4j4z3S/tMftN+E/hFp5+0eEfAs8XivxpNGf3TXK/8eFgf9pmzK6f3f8AaWm+OPiZ4j8OeMZ/jMuo+Fbf"
    "4cS32neErS311JEuLmze9WK7u4J/NWKJXd3f50dXitIn+Wuf1bwzb/s+fA/xF4A0DxLdHxbqX+leM/HSokt3NqN58qRRbnXd"
    "dXDuixJvXyom81nX5d/pGg+NNZh+Ingz4V6z8MLDxWND0G1l1jxFoqWqado11KrRJFFbyvvij8pJunzbCu1NrUAdF+zH4d8C"
    "aTpviKT4baDb6b4ON6sFlq0NzLcLqqxJ87wvIzZtkd3ii2ts+SXaNuN3A/sjv/wt343fGv42f67StQ1GPwp4cmB+VrCx+WWV"
    "P9iWX5v+A10X7ZnxQ1Xwr4F0/wCHngYq/wASvH8v9iaHbxcG1ibi4vG/upFF/F/CxWvXPg/8MtJ+DXwz8M+CNCTGmaJYraxy"
    "Mu1pWHLyt/tu5Z2/2noA7aiiigAooooAKKKKACiiigAooooAKKKKACvjPxJDP+wv8XNS8YW0Er/AbxpfpLr8NujMvhjVJcL9"
    "tVFH/HtKdqv/AHWx/sI32ZWVrmh6f4l0W90fVbSHUNNvoWt7mzuIw8U8bLtZGU/wkGgD5b+MX7PfhuRvHPxH1/Xtb1/4V/YT"
    "4gf4d+F3eKy1GWOHzZbiQrLidpfv7V8pG4Z953M2L4c/aC0T4b+OtA0Pwl4O1XS9V15I7i++GWlW0V0kFm0W5dXspbV3gWPa"
    "vzJuXzeqoku/zZ4Y/FX7AdxMiQah4z/Z2kk3Ky77rUvB+772V+9PZ/8Ajyf+h53xkt/h5pvh3wpafBbwXoaeJviNqytonjvS"
    "IPsun6Vc7flupL2L/lrt3qturfvSXRl+dlYA7jxv8C4fiBrifGX4DeKrTwn4/nUJdT7HOm68qn5rfUrf7yyL8ybtvmx+m5V2"
    "aXgf9sLSodeg8HfGHSZfhD49fCR22szL/ZmoYzl7K9/1ci9PlYq2W2/NXcfBX9nvRPgvHNc22o6vrPiTUS8uua1f3kpbVrpy"
    "rNcSxbvL3fKFXavyp8td14x8C6B8QtBn0TxLomn6/pE3MtlqVqs8TY/2W/i/2qANxJFkRWVtyt91lrg/jf8AB3Qvjv8ADTWf"
    "BfiBWWy1GP8Ad3EP+ttp1+aKeP8A20b5vwryKL9jXUPh3k/Bj4p+J/hpAv8AqtCu2/tvRUHfba3PzR/8Blqwus/tV+DVIuvD"
    "Xw5+JNsvCSaZqd1o13J7ukqSxKf91qAJf2TPi/r/AIht9c+FnxHkJ+K3gR1tdSkZsf2vZ/8ALvqEX95ZV27v9r723cFrjfjh"
    "e3H7WnxpX4GaFdTL8P8Aw5LFqPxC1O3JEdy33rfSEb+85XdJt+7t67lZa81/aCsv2i/Hnijwv4+8F/Ai98D/ABO0AS2kWs2/"
    "ijTL+2urOVW8y3mTejOu75k3LhWzXW/s9/8AC6/gf8NrLwzon7OU9zqU0r32sa9r3jewjl1G/l+ae4l8vzW+Zxx/shaAPtHT"
    "9PttJsYLOzhjtbW3jWGGCNdqRqoCqqqOi1W8ReJNJ8J6LcatrmqWmjaXbJvnvdQnWCGJfVnb5a8Ck0n9qjx5GYrnXvh78KbK"
    "UcTaRZ3GuajF/wB//Kg/8dapdF/Yj8GXms2mu/ErV9e+MviCB/Njl8Y3n2iyt3P3vIslCwIn+zsagDn9U/aV8ZfH65l0D9nv"
    "SRJphbybz4n+ILd4tIs/7xs4m+a8lXn/AKZbtu7cjVynwxbwZ8L/ABla+BPBfiGa78eeNmmn1f4r+IYHlbWZoivmxWUzL5U8"
    "y7/kiVvKi28+a6ujesftCeA4rjwPq9x4hsvFnjbwzFFstvBPgomwXykTdtl8uVJZfuH5d+37ipFu+94pb+C9P+J3wJ03Ttc0"
    "7xZ8SNS0PR5RFrWrXM+g6fpFzudoJYpbxbdnniDxp9q8p2VIH+6zPE4A+++D+t/E/wAdSeDb3wLfQ+HdH8Surf8ACR2KzaRP"
    "pP8Ay3vfPdt9zqN6x+WVfnh/vRfP5v0Z4k1r4dfsjfCTUtaltbfwz4a0/dK0VooaW7nb7qru+aWZ+F+Y5P0FfPHhH9oPXvgn"
    "eRar44+Klx8Q/AMehR6bYzJokKXXiLXBKq7NJWL97dLt+VpW3R73T5/v7O5+GPwT8XfG3x5pnxb+N9oNOfTW8/wn8PVl8630"
    "Xut1df8APW8/RP8Ae4QAv/szfC/xL4h8Yat8ePilaiz8beIoPsui6FNlh4b0nO5Lcf8ATZ87pW/9A+Za+nqKKACiiigAoooo"
    "AKKKKACiiigAooooAKKKKACiiigCJ41kRlZdyt95Wr5R+IX7Hd/4VuNT1z4HahZ+FpNQuBdap4H1YM/hzVZUdGV1jX5rOfci"
    "Mstvt5Rfu9a+s6KAPmvwn+2hoOma1B4Y+Lmi3vwZ8YSfKkPiBlbTLxv71rfr+6kX/e2tX0bbzxXkMc0MqywyLuV1bcrLWX4r"
    "8I6J460O40bxDpFlrulXC7ZbLUbdJ4X/AN5GHNeBN+xPYeCZ3u/g5498UfCGZjvGm6fc/wBoaOz/AN97C53L/wB8MtAH0xRX"
    "zMmr/tT+Ady3Wh+APixYxfKj6fez6DqE3u6yrLAG/wB1ql/4at8a6Cu3xX+zp8SrC5P3v+EfSy1qJf8AgcVx/wCy0AfSlFfN"
    "0f7cHh6OPN78Nvizp0n9268DXob/AMdU0sn7bOlzLnTvhN8YNXyv/Lp4IulH/fUuxaAPpCivmlf2kvix4jUHwl+zd4rkRm2C"
    "bxdq9hoyr/tMm+V/++VqNvDH7UXxE+TVvGHgv4Taa/IXwzp8us6gF/uPLc7IlP8AtKjUAe++LvGmg+AdDn1jxHrNjoGkwf62"
    "91K5WCJf+BNxXyX4w+KKftUa9Zj4O/C6w8bCzHlw/EXxvYPFoWnHd9+1ilXfcyr/ALCr/D8xWvRvC/7Evw7sdbt/EHjNtZ+K"
    "/iaE5j1bx5fvqPl+0Vv/AKhF/wCAfLxX0JDBHawpFGixxIu1VUYVVoA8N+Df7KulfD/xFJ428V61efEb4o3EflzeKtaRc2yf"
    "88rK3UbLWLr8q/N8zfNzXvNFFABRRRQAUUUUAFFFFABRRRQAUUUUAFFFFABRRXzH8bPj9441L4vH4MfB7TtPPjOPTl1XWPEm"
    "vBvsGi2rPsXbEvzTztlSq/d+ZPvfPsAPpyolmjdiqurOv8O6vl61/YWsvGSm6+LvxL8b/FC8l5ntJdVl0vSv+2dpasmxf+B1"
    "wv7Qn7G37OvwP+Cvi3x3B4Em0m70Cxa6tLzR9avILtJ/uxFZfN+9vZfvbutAH3BRXwz+zr+098UPAvwT8J6b8RPgv8UvFuvQ"
    "2ZafXNMsbe9+2Rs7NE/M6tuETIvzfM22vSYf27NNh/5CXwX+NOix95L7wTKVX8YnegD6dor5i/4eGfCq3H/EwtPGejH/AKfv"
    "CV+v/oMTU1v+Ck37PMfF146urF/7t14e1OI/rb0AfT9FfNUP/BRz9nCYfL8UdNB/2rS6X/2lV6P/AIKCfs7yL8vxW0If73mr"
    "/wCyUAfQ9FfP/wDw3t+z3/0Vjw7/AN/2/wDiaVv2+P2e1Xd/wtjw/wD9/X/+JoA9/or51n/4KDfs7xdfitoh/wB1ZW/9krOn"
    "/wCCkX7OESf8lPs3/wBmHT712/8AHYaAPpuivmJf+CkPwCmx9k8WalqJb7v2Twzqj7v/ACXob/goF8ObjA0zw/8AEDWS33fs"
    "Pg+/bd/30i0AfTtFfMX/AA3JFdSbdP8AgN8btS/uuvg0xJ/31LKtfOn7YHx68WePJfh2fEfw6+Ifw/8Ag5FrKweLmvbiPTXv"
    "4p3VIkZoJXbYuHdk/iz/AA/K1AH6RrNG7FVdWdf4d1S18wv/AME3f2dFjVYPhytpNHylxa6vfxSq397etxuqhffsr/Ef4Wxv"
    "qHwV+MWvQzRfOvhbx9O2r6TPjpErt+/t16fMrMaAPq2ivEf2Y/2hJfjnoviG31nQ38L+NvCupNoviHRzL5sUN0n8cMn8cT8l"
    "f68M3t1ABRRRQAUUUUAFFFFABRRRQAV8B/tyfCHRT+1V8EfHOrXmr6NoviKdvBuq6loOotYXME773snEq9PnZ93+ylffleDf"
    "tr/CC8+N37Nni7QdGRpPElrGmqaN5PMv2y2YSoqf7T7Wj/4HQB82+MLKy/ZG/bJ+D39qfFrxhL4N1bS9Yk1WXxx4okurT91b"
    "t5K5l2r99k+X+9sqr+2F+1N8Ov2ovCfhP4PfD3W7rX5fF/izS9Ou7qHTriG1NqtwGl2SyoquVdYj8m6sOy+L3hX9q79sT9m6"
    "21fRkv8AXNK0TVF8UeG9b0h1GnX4tN+x4p02/LKm5G5/g+61faXxu/Zj8FfHqPw3/b7arptz4blln0q70G/exmtWcIHKsn+4"
    "tAHrMcSwxrHEqoirtVV/hr5m8H/Ff4k3v7R2pya14b12x+GOpR3+m6L9ptIkhEtnsdbhvm89PP2Xv+tRF2pb7d2+hP2HXs23"
    "ad8fvjbYn+GNvF3nov8AwGWJqlX9lb4m6b/yCP2mvHUR7DUrGwvf/Q4qAM/SP2mJv+GbvAHiW78T2L+Jhd+HLfxVcXYiiWzN"
    "zcW/23zl2qsX7r7R/d21F8Z/2kLqw+K3gnS/C3jfRrLwvrGkSah/aH2zTvKun+1JEnly3Mqq/wDH8kWWrRb4A/tE27f6H+1N"
    "M6D7q33gHTZf/HlZKRvg/wDtOwLtHx98N6jt+79r8CxJu/75uKALHh/x98RfGP7Q/iTRrTSLC8+Hui6+2l6hdXFlCqQQ/wBl"
    "W8+5Z/tHmvP586Ls8jZsf7/y15boHxu+JV1+zvrvxH1nwZ4fS0/sPS9Q0e8u9Gjt0lurm42yxeV9tbzYlieFlZzB8zf98+kw"
    "+Af2rLLcYfid8Ornc25vO8Lzpu/75lqvP4H/AGqv7NOnv4k+EF5Ybdn2SfQL0RMv93Z5u2gClrfjjXfDPwbsfG9t4T8KeJp1"
    "1Mabfad9gs7Vm891t7VkeC9uol23Etvu/et+6dvlVl+b3zwv4T0j+y7aHUNN8PXeu2kUUOpPptkqQpceWrsFU7mRTuDKjNna"
    "614Vb+Af2pLOw+w2uo/BO2svN837NDoWoLFv3bt+3zfvbhUtl4T/AGsdPutQubXXfg3DPfzrcXTrpep/vJViSLd/rf7kSL/w"
    "GgDjvAv7R2v+IPCvjzUL/wAMaB4Zv9C8JanrGkQXGnH/AInU9rLcRNdJ8/ywRPBEjRff/eo/yqyVeh/aB8c3nwd0nWNIvbPU"
    "vEd94l07SHi0q10yWRbedPuIsWoSxK+77vmyxf7tdGPh/wDtSyrEreKfhHbLGrJH5Ph68farfe+9L/F3p1v8KP2noV2wfE/4"
    "eaUm7dtsfCD/APs0tAHKa1+018R7j4L/AA9m8OWJ174iapeajdahp+naU8sy2VhLKstvcQfN9nuGdrO3l/hSWWXZ8qrXrXhf"
    "9oCw1H4uSWN3rDJ4c8QeHtC1XwzDNaFGllupb1JV3Ku7d+6tdyP9zd71zb/Bn9pu8ZfN/aN0axXv9l8A27/+h3FJH+zv8frj"
    "d9v/AGptTdT2svBOmW//AMVQBJ8FPiZ8S9Y+NHidfE/hzXoPh9rlxdf8I5e3FlEsNn9kleLYfKbzUS4iVJVNwifMjqu7etbv"
    "7dngj/hYf7IvxR0lIzNIukPqESp95ntXW5UL+MVYDfskeOtUX/ib/tJ/EyT/ALBj2dl/6DBVWb9g3S9Uhlg1v40/GbxBayps"
    "ks9R8ZMYZF/iyiRLQBk/BD/got8G/FXgnwjb+JfGS+GfEtzplv8Aak162ns4nn8pfNZbh08pl37vm314H+y/8LbH4jfsm6p8"
    "XvHnxU+KJitG1S8vrXS/F9xBazQWzy4wn+4v96vvLwX8B/BPgn4U6b8OLbQrbU/CGnxvFHp2roL1CrOztu83du+Z2/OvzC+D"
    "XjUfFD9jnQv2bvBP9oXHjPX/ABPJb6+lnYy+Xo+lG8eWW4lk2eUo2oi7d3dqAPsn/gmR8K5Phz+y3o+rXcLRav4xuZfEFz5z"
    "F32S/Lbjc3zN+5SN+e7tX1zWZoui2Xh/RtP0mwgS3sLCFLa2hT7scaKFRf8AvnFadABRRRQAUUUUAFFFFABRRRQAUUUUAQeT"
    "E0izFFLqu1Xx81T0UUAFFFFABRRRQAUUUUAFFFFABRRRQAUUUUAFFFFABUFvDFDHthRY0znai7RRRQBPRRRQAUUUUAFFFFAH"
    "/9k="
)


def _load_logo_bytes():
    try:
        return base64.b64decode(_LOGO_BASE64)
    except Exception:
        return None



def _grade_sort_key(g):
    return GRADE_ORDER.index(g) if g in GRADE_ORDER else len(GRADE_ORDER)


def effective_grade(student):
    """Grade Awarded if present, else 'F' if the overall result is Fail."""
    if student["grade_awarded"]:
        return student["grade_awarded"]
    if student["status"].lower() == "fail":
        return "F"
    return None


# --------------------------------------------------------------------------
# ---------------------------  PDF PARSING  -------------------------------
# --------------------------------------------------------------------------

def _to_float(val):
    """Safely convert a mark-sheet cell to float, else return None."""
    if val is None:
        return None
    val = str(val).strip()
    if val in ("", "*", "-", "--", "Ab"):
        return None
    try:
        return float(val)
    except ValueError:
        return None


def parse_marklist_pdf(file_bytes):
    """
    Parse the uploaded PDF (bytes) and return:
        students      : list of dicts, one per student
        course_names  : dict {course_code: course_title}
    """
    students = []
    course_names = {}

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""

            m_id = re.search(r"APAAR ID\s*:\s*([0-9A-Za-z]+)", text)
            m_name = re.search(r"Name of Student\s*:\s*(.+)", text)
            m_prog = re.search(r"Name of Programme\s*:\s*(.+)", text)

            if not m_id or not m_name:
                continue

            apaar = m_id.group(1).strip()
            name = m_name.group(1).strip()
            programme = m_prog.group(1).strip() if m_prog else ""

            tables = page.extract_tables()
            if not tables:
                continue
            main_table = tables[0]

            courses = []
            total_row = None
            result_row = None

            for row in main_table:
                first = (row[0] or "").strip() if row[0] else ""

                if first == "Total":
                    total_row = row
                    continue
                if first.startswith("Result:"):
                    result_row = row
                    continue
                if row[1] is None or first in ("", "Course Code"):
                    continue

                code = first
                title = (row[1] or "").replace("\n", " ").strip()
                if not code or not title:
                    continue

                credits = row[2]
                course_total = _to_float(row[13])
                gp = _to_float(row[14])
                cp = _to_float(row[15])
                grade = (row[16] or "").strip()

                course_names.setdefault(code, title)
                courses.append(
                    {
                        "code": code,
                        "title": title,
                        "credits": credits,
                        "course_total": course_total,
                        "gp": gp,
                        "cp": cp,
                        "grade": grade,
                    }
                )

            total_credits = ""
            total_cp = None
            if total_row:
                total_credits = (total_row[2] or "").strip()
                for idx in (15, 14, 13):
                    val = total_row[idx]
                    if val:
                        total_cp = _to_float(val)
                        break

            status, sgpa, grade_awarded = "Unknown", None, ""
            if result_row:
                joined = " ".join([c for c in result_row if c])
                m_status = re.search(r"Result\s*:\s*(Pass|Fail)", joined, re.I)
                m_sgpa = re.search(r"SGPA\s*:\s*([\d.]+|--)", joined)
                m_grade = re.search(r"Grade awarded\s*:\s*([A-Za-z+\-]+|--)", joined)
                if m_status:
                    status = m_status.group(1).capitalize()
                if m_sgpa and m_sgpa.group(1) != "--":
                    sgpa = float(m_sgpa.group(1))
                if m_grade and m_grade.group(1) != "--":
                    grade_awarded = m_grade.group(1)

            total_marks = sum(c["course_total"] for c in courses if c["course_total"])
            total_marks = round(total_marks, 2)

            students.append(
                {
                    "name": name,
                    "apaar": apaar,
                    "programme": programme,
                    "courses": courses,
                    "total_credits": total_credits,
                    "total_cp": total_cp,
                    "sgpa": sgpa,
                    "grade_awarded": grade_awarded,
                    "status": status,
                    "total_marks": total_marks,
                }
            )

    return students, course_names


# --------------------------------------------------------------------------
# ---------------------------  ANALYSIS  -----------------------------------
# --------------------------------------------------------------------------

def build_student_table(students):
    rows = []
    for i, s in enumerate(students, start=1):
        rows.append(
            {
                "Sl.No": i,
                "Name": s["name"],
                "APAAR ID": s["apaar"],
                "Total Credits": s["total_credits"] if s["total_credits"] else "--",
                "SGPA": s["sgpa"] if s["sgpa"] is not None else "--",
                "Grade Awarded": s["grade_awarded"] if s["grade_awarded"] else "--",
                "Status": s["status"],
            }
        )
    return pd.DataFrame(rows)


def build_result_summary(students):
    appeared = len(students)
    passed = sum(1 for s in students if s["status"].lower() == "pass")
    failed = sum(1 for s in students if s["status"].lower() == "fail")
    other = appeared - passed - failed
    pass_pct = round((passed / appeared) * 100, 2) if appeared else 0.0

    summary_df = pd.DataFrame(
        [
            {"Category": "Appeared", "Count": appeared},
            {"Category": "Passed", "Count": passed},
            {"Category": "Failed", "Count": failed},
        ]
        + ([{"Category": "Result Not Available", "Count": other}] if other else [])
    )
    return summary_df, pass_pct


def build_top_rankers(students, top_n=3):
    ranked = [s for s in students if s["sgpa"] is not None]
    if not ranked:
        return pd.DataFrame()
    ranked.sort(key=lambda s: s["sgpa"], reverse=True)

    distinct_sgpas = sorted({s["sgpa"] for s in ranked}, reverse=True)
    cutoff_values = distinct_sgpas[:top_n]

    rows = []
    for pos, val in enumerate(cutoff_values, start=1):
        holders = [s for s in ranked if s["sgpa"] == val]
        for s in holders:
            rows.append(
                {
                    "Rank": pos,
                    "Name": s["name"],
                    "APAAR ID": s["apaar"],
                    "SGPA": s["sgpa"],
                    "Grade Awarded": s["grade_awarded"],
                }
            )
    return pd.DataFrame(rows)


def build_subject_failure_table(students, course_names):
    fail_counts = {code: 0 for code in course_names}
    appeared_counts = {code: 0 for code in course_names}

    for s in students:
        for c in s["courses"]:
            code = c["code"]
            appeared_counts[code] = appeared_counts.get(code, 0) + 1
            if c["grade"].upper() == "F":
                fail_counts[code] = fail_counts.get(code, 0) + 1

    rows = []
    for code in course_names:
        rows.append(
            {
                "Course Code": code,
                "Course Title": course_names[code],
                "Appeared": appeared_counts.get(code, 0),
                "Failed": fail_counts.get(code, 0),
            }
        )
    df = pd.DataFrame(rows).sort_values("Failed", ascending=False).reset_index(drop=True)
    return df


def build_grade_distribution(students):
    """Grade Awarded vs. No. of Students, ordered O, A+, A, B+, B, C, P, F."""
    counts = Counter()
    for s in students:
        g = effective_grade(s)
        if g:
            counts[g] += 1
    ordered_grades = sorted(counts.keys(), key=_grade_sort_key)
    rows = [{"Grade Awarded": g, "No. of Students": counts[g]} for g in ordered_grades]
    return pd.DataFrame(rows)


def build_grade_wise_student_list(students):
    """Sl.No, Name, Grade — rows grouped together by grade, in hierarchy order."""
    groups = {}
    for s in students:
        g = effective_grade(s) or "NA"
        groups.setdefault(g, []).append(s)

    ordered_grades = sorted(groups.keys(), key=_grade_sort_key)
    rows = []
    sl_no = 1
    for g in ordered_grades:
        for s in sorted(groups[g], key=lambda x: x["name"]):
            rows.append({"Sl.No": sl_no, "Name": s["name"], "Grade": g})
            sl_no += 1
    return pd.DataFrame(rows), ordered_grades, groups


def build_summary_paragraphs(
    exam_details,
    exam_month_year,
    programme_name,
    summary_df,
    pass_pct,
    top_df,
    fail_df,
    avg_sgpa,
    highest_sgpa,
    lowest_sgpa,
):
    """Return a list of paragraph strings summarising the whole report."""
    def _count(cat):
        rows = summary_df.loc[summary_df["Category"] == cat, "Count"]
        return int(rows.iloc[0]) if not rows.empty else 0

    appeared = _count("Appeared")
    passed = _count("Passed")
    failed = _count("Failed")

    paras = []

    exam_line = exam_details
    if exam_month_year:
        exam_line = f"{exam_details} ({exam_month_year})" if exam_details else exam_month_year

    paras.append(
        f"A total of {appeared} student(s) appeared for the {exam_line} "
        f"held for the {programme_name} programme. Of these, {passed} "
        f"student(s) passed and {failed} student(s) failed, resulting in "
        f"an overall pass percentage of {pass_pct}%."
    )

    if avg_sgpa is not None:
        paras.append(
            f"Among the students who passed, the average SGPA secured was "
            f"{avg_sgpa}, with the highest SGPA recorded at {highest_sgpa} "
            f"and the lowest passing SGPA at {lowest_sgpa}."
        )

    if top_df is not None and not top_df.empty:
        top_rank1 = top_df[top_df["Rank"] == 1]
        names = ", ".join(top_rank1["Name"].tolist())
        sgpa_val = top_rank1["SGPA"].iloc[0]
        verb = "secured" if len(top_rank1) == 1 else "jointly secured"
        paras.append(f"{names} {verb} the top rank with an SGPA of {sgpa_val}.")

    if fail_df is not None and not fail_df.empty and fail_df["Failed"].sum() > 0:
        top_fail = fail_df.iloc[0]
        paras.append(
            f"The highest number of failures was recorded in "
            f"{top_fail['Course Code']} ({top_fail['Course Title']}), where "
            f"{top_fail['Failed']} student(s) failed — indicating this "
            f"course may need additional academic support."
        )
    else:
        paras.append("No student failed in any individual course.")

    return paras


# --------------------------------------------------------------------------
# ---------------------------  DOCX BUILDING  ------------------------------
# --------------------------------------------------------------------------

def _set_cell_shading(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _add_bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_heading(doc, text, size=13):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p


def _add_table_from_df(doc, df, header_color="4472C4", header_font_color="FFFFFF"):
    if df is None or df.empty:
        doc.add_paragraph("(No data)")
        return
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr_cells[i].text = str(col)
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(header_font_color)
        _set_cell_shading(hdr_cells[i], header_color)

    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(df.columns):
            cells[i].text = "" if pd.isna(row[col]) else str(row[col])
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def _fig_to_docx(doc, fig, width_inches=5.5):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=160, bbox_inches="tight")
    buf.seek(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(buf, width=Inches(width_inches))


def _add_signature_block(doc):
    """Three signature slots — Faculty Advisor, HOD, Principal — in one row."""
    # Leave some blank vertical space for the actual pen-and-ink signature.
    doc.add_paragraph()
    doc.add_paragraph()

    labels = ["Faculty Advisor", "Head of the Department", "Principal"]

    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_width = Inches(2.1)
    for row in table.rows:
        for cell in row.cells:
            cell.width = col_width
    for col in table.columns:
        col.width = col_width

    for i, label in enumerate(labels):
        line_p = table.cell(0, i).paragraphs[0]
        line_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line_p.add_run("_" * 24)

        label_p = table.cell(1, i).paragraphs[0]
        label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_run = label_p.add_run(label)
        label_run.bold = True
        label_run.font.size = Pt(11)


def generate_docx_report(
    college_name,
    exam_details,
    exam_month_year,
    programme_name,
    logo_bytes,
    student_df,
    summary_df,
    pass_pct,
    top_df,
    fail_df,
    legend_df,
    grade_dist_df,
    grade_wise_df,
    avg_sgpa=None,
    highest_sgpa=None,
    lowest_sgpa=None,
    fig_summary=None,
    fig_fail=None,
    fig_grade=None,
):
    doc = Document()

    # ---------------- Letterhead / Header ----------------
    if logo_bytes:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(io.BytesIO(logo_bytes), width=Inches(0.9))

    if college_name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(college_name)
        run.bold = True
        run.font.size = Pt(16)

    if exam_details or exam_month_year:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if exam_details and exam_month_year:
            line = f"{exam_details}, {exam_month_year}"
        else:
            line = exam_details or exam_month_year
        run = p.add_run(line)
        run.font.size = Pt(12)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Result Analysis Report")
    run.bold = True
    run.underline = True
    run.font.size = Pt(14)

    if programme_name:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run(f"Programme: {programme_name}")
        run2.font.size = Pt(12)

    hr_p = doc.add_paragraph()
    _add_bottom_border(hr_p)

    # ---------------- 1. Consolidated Result Table ----------------
    _add_heading(doc, "1. Consolidated Result Table")
    _add_table_from_df(doc, student_df)

    # ---------------- 2. Result Summary ----------------
    _add_heading(doc, "2. Result Summary")
    _add_table_from_df(doc, summary_df)
    p = doc.add_paragraph()
    run = p.add_run(f"Pass Percentage: {pass_pct}%")
    run.bold = True
    if fig_summary is not None:
        _fig_to_docx(doc, fig_summary, width_inches=4.5)

    # ---------------- 3. Top Rank Holders ----------------
    _add_heading(doc, "3. Top Rank Holders")
    _add_table_from_df(doc, top_df)

    # ---------------- 4. Subject-wise Failure Analysis ----------------
    _add_heading(doc, "4. Subject-wise Failure Analysis")
    _add_table_from_df(doc, fail_df)
    if not fail_df.empty and fail_df["Failed"].sum() > 0:
        top_row = fail_df.iloc[0]
        p = doc.add_paragraph()
        run = p.add_run(
            f"Highest number of failures: {top_row['Course Code']} — "
            f"{top_row['Course Title']} ({top_row['Failed']} student(s))"
        )
        run.italic = True
    if fig_fail is not None:
        _fig_to_docx(doc, fig_fail, width_inches=5.5)

    # ---------------- 5. Course Code / Full Name Reference ----------------
    _add_heading(doc, "5. Course Code / Full Name Reference")
    _add_table_from_df(doc, legend_df)

    # ---------------- 6. Grade Distribution ----------------
    _add_heading(doc, "6. Grade Distribution")
    sgpa_stats_df = pd.DataFrame(
        [
            {"Metric": "Average SGPA (Passed Students)", "Value": avg_sgpa if avg_sgpa is not None else "--"},
            {"Metric": "Highest SGPA", "Value": highest_sgpa if highest_sgpa is not None else "--"},
            {"Metric": "Lowest SGPA (among passed)", "Value": lowest_sgpa if lowest_sgpa is not None else "--"},
        ]
    )
    _add_table_from_df(doc, sgpa_stats_df)
    doc.add_paragraph()
    _add_table_from_df(doc, grade_dist_df)
    if fig_grade is not None:
        _fig_to_docx(doc, fig_grade, width_inches=4.5)

    # ---------------- 7. Grade-wise List of Students ----------------
    _add_heading(doc, "7. Grade-wise List of Students")
    _add_table_from_df(doc, grade_wise_df)

    # ---------------- 8. Summary ----------------
    _add_heading(doc, "8. Summary")
    summary_paragraphs = build_summary_paragraphs(
        exam_details=exam_details,
        exam_month_year=exam_month_year,
        programme_name=programme_name,
        summary_df=summary_df,
        pass_pct=pass_pct,
        top_df=top_df,
        fail_df=fail_df,
        avg_sgpa=avg_sgpa,
        highest_sgpa=highest_sgpa,
        lowest_sgpa=lowest_sgpa,
    )
    for para_text in summary_paragraphs:
        p = doc.add_paragraph(para_text)
        p.paragraph_format.space_after = Pt(8)

    # ---------------- Signatures ----------------
    _add_signature_block(doc)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# --------------------------------------------------------------------------
# ---------------------------  STREAMLIT UI  -------------------------------
# --------------------------------------------------------------------------

st.set_page_config(page_title="Result Analysis Report", layout="wide")

st.title("📊 Result Analysis Report Generator")
st.caption(
    "Upload a University mark-cum-grade-statement PDF (one page per student) "
    "and generate a consolidated result analysis report."
)

college_name = COLLEGE_NAME
logo_bytes = _load_logo_bytes()

with st.sidebar:
    st.header("Report Details")
    st.markdown(f"**Institution:** {COLLEGE_NAME}")
    exam_details = st.text_input(
        "Examination Details", "Fourth Semester FYUGP Examination"
    )
    programme_name = st.text_input("Programme / Major", "Physics")
    exam_month_year = st.text_input("Month & Year", "May 2026")

    st.markdown("---")
    uploaded_pdf = st.file_uploader("Upload Mark-list PDF", type=["pdf"])
    generate = st.button("Generate Report", type="primary", use_container_width=True)

if generate and uploaded_pdf is not None:
    with st.spinner("Reading and parsing the PDF..."):
        students, course_names = parse_marklist_pdf(uploaded_pdf.read())

    if not students:
        st.error(
            "Could not find any student mark statements in this PDF. "
            "Please check the file and try again."
        )
        st.stop()

    st.session_state["students"] = students
    st.session_state["course_names"] = course_names

if "students" in st.session_state:
    students = st.session_state["students"]
    course_names = st.session_state["course_names"]

    # ---------------- Report Header ----------------
    st.markdown("---")
    header_html = "<div style='text-align:center'>"
    if logo_bytes:
        import base64
        b64 = base64.b64encode(logo_bytes).decode()
        header_html += f"<img src='data:image/png;base64,{b64}' style='height:70px'><br>"
    if college_name:
        header_html += f"<h2 style='margin-bottom:0'>{college_name}</h2>"
    header_html += f"<h4 style='margin-top:4px'>{exam_details} &nbsp;&mdash;&nbsp; {exam_month_year}</h4>"
    header_html += f"<h3 style='margin-top:4px;text-decoration:underline'>Result Analysis Report</h3>"
    header_html += f"<p style='font-size:16px'><b>Programme:</b> {programme_name}</p>"
    header_html += "</div>"
    st.markdown(header_html, unsafe_allow_html=True)
    st.markdown("---")

    # ---------------- 1. Consolidated Student Table ----------------
    st.subheader("1. Consolidated Result Table")
    student_df = build_student_table(students)
    st.dataframe(student_df, use_container_width=True, hide_index=True)

    # ---------------- 2. Result Summary ----------------
    st.subheader("2. Result Summary")
    summary_df, pass_pct = build_result_summary(students)
    col1, col2 = st.columns([1, 1.4])
    fig_summary = None
    with col1:
        st.dataframe(summary_df, use_container_width=True, hide_index=True)
        st.metric("Pass Percentage", f"{pass_pct}%")
    with col2:
        fig_summary, ax = plt.subplots(figsize=(4, 3))
        colors = {"Appeared": "#4C72B0", "Passed": "#55A868", "Failed": "#C44E52",
                  "Result Not Available": "#8172B2"}
        bar_colors = [colors.get(c, "#999999") for c in summary_df["Category"]]
        ax.bar(summary_df["Category"], summary_df["Count"], color=bar_colors)
        ax.set_ylabel("No. of Students")
        ax.set_title("Appeared / Passed / Failed")
        for i, v in enumerate(summary_df["Count"]):
            ax.text(i, v + 0.1, str(v), ha="center", fontweight="bold")
        plt.xticks(rotation=15)
        st.pyplot(fig_summary, use_container_width=True)

    # ---------------- 3. Top Rank Holders ----------------
    st.subheader("3. Top Rank Holders")
    top_df = build_top_rankers(students, top_n=3)
    if top_df.empty:
        st.info("No student has a valid SGPA to rank (all results unavailable).")
    else:
        st.dataframe(top_df, use_container_width=True, hide_index=True)

    # ---------------- 4. Subject-wise Failure Analysis ----------------
    st.subheader("4. Subject-wise Failure Analysis")
    fail_df = build_subject_failure_table(students, course_names)
    st.dataframe(fail_df, use_container_width=True, hide_index=True)

    fig_fail = None
    if fail_df["Failed"].sum() > 0:
        top_fail_row = fail_df.iloc[0]
        st.warning(
            f"📌 Highest number of failures is in **{top_fail_row['Course Code']} "
            f"— {top_fail_row['Course Title']}** with **{top_fail_row['Failed']}** "
            "student(s) failing."
        )
        fig_fail, ax2 = plt.subplots(figsize=(8, 4))
        plot_df = fail_df[fail_df["Failed"] > 0]
        ax2.bar(plot_df["Course Code"], plot_df["Failed"], color="#C44E52")
        ax2.set_ylabel("No. of Students Failed")
        ax2.set_title("Failures by Course")
        for i, v in enumerate(plot_df["Failed"]):
            ax2.text(i, v + 0.05, str(v), ha="center", fontweight="bold")
        plt.xticks(rotation=30, ha="right")
        st.pyplot(fig_fail, use_container_width=True)
    else:
        st.success("No student failed in any subject. 🎉")

    # ---------------- 5. Course Code -> Full Name Legend ----------------
    st.subheader("5. Course Code / Full Name Reference")
    legend_df = pd.DataFrame(
        [{"Course Code": k, "Course Full Name": v} for k, v in course_names.items()]
    ).sort_values("Course Code").reset_index(drop=True)
    st.dataframe(legend_df, use_container_width=True, hide_index=True)

    # ---------------- 6. Grade Distribution ----------------
    st.subheader("6. Grade Distribution")
    grade_dist_df = build_grade_distribution(students)
    valid_sgpas = [s["sgpa"] for s in students if s["sgpa"] is not None]
    avg_sgpa = round(sum(valid_sgpas) / len(valid_sgpas), 2) if valid_sgpas else None
    highest_sgpa = max(valid_sgpas) if valid_sgpas else None
    lowest_sgpa = min(valid_sgpas) if valid_sgpas else None

    c1, c2, c3 = st.columns(3)
    with c1:
        st.metric("Average SGPA (Passed Students)", avg_sgpa if avg_sgpa is not None else "--")
    with c2:
        st.metric("Highest SGPA", highest_sgpa if highest_sgpa is not None else "--")
    with c3:
        st.metric("Lowest SGPA (among passed)", lowest_sgpa if lowest_sgpa is not None else "--")

    fig_grade = None
    if not grade_dist_df.empty:
        colA, colB = st.columns([1, 1.4])
        with colA:
            st.dataframe(grade_dist_df, use_container_width=True, hide_index=True)
        with colB:
            fig_grade, ax3 = plt.subplots(figsize=(4, 3))
            ax3.bar(grade_dist_df["Grade Awarded"], grade_dist_df["No. of Students"], color="#4C72B0")
            ax3.set_title("Grade Distribution")
            ax3.set_ylabel("No. of Students")
            st.pyplot(fig_grade, use_container_width=True)

    # ---------------- 7. Grade-wise List of Students ----------------
    st.subheader("7. Grade-wise List of Students")
    grade_wise_df, ordered_grades, groups = build_grade_wise_student_list(students)
    st.dataframe(grade_wise_df, use_container_width=True, hide_index=True)

    # ---------------- 8. Summary ----------------
    st.subheader("8. Summary")
    summary_paragraphs = build_summary_paragraphs(
        exam_details=exam_details,
        exam_month_year=exam_month_year,
        programme_name=programme_name,
        summary_df=summary_df,
        pass_pct=pass_pct,
        top_df=top_df,
        fail_df=fail_df,
        avg_sgpa=avg_sgpa,
        highest_sgpa=highest_sgpa,
        lowest_sgpa=lowest_sgpa,
    )
    for para_text in summary_paragraphs:
        st.write(para_text)

    st.markdown("#####")
    sig1, sig2, sig3 = st.columns(3)
    for col, label in zip((sig1, sig2, sig3), ("Faculty Advisor", "Head of the Department", "Principal")):
        with col:
            st.markdown("&nbsp;")
            st.markdown("____________________")
            st.markdown(f"**{label}**")

    st.markdown("---")
    st.caption(
        "Note: 'Total Credits' and 'SGPA' are taken as printed on the "
        "official mark statement (shown as '--' where the university's "
        "statement itself shows '--', e.g. when the result is Fail). "
        "In the Grade Distribution and Grade-wise List, students whose "
        "overall result is Fail are grouped under grade 'F'."
    )

    # ---------------- DOCX download ----------------
    st.markdown("---")
    docx_bytes = generate_docx_report(
        college_name=college_name,
        exam_details=exam_details,
        exam_month_year=exam_month_year,
        programme_name=programme_name,
        logo_bytes=logo_bytes,
        student_df=student_df,
        summary_df=summary_df,
        pass_pct=pass_pct,
        top_df=top_df,
        fail_df=fail_df,
        legend_df=legend_df,
        grade_dist_df=grade_dist_df,
        grade_wise_df=grade_wise_df,
        avg_sgpa=avg_sgpa,
        highest_sgpa=highest_sgpa,
        lowest_sgpa=lowest_sgpa,
        fig_summary=fig_summary,
        fig_fail=fig_fail,
        fig_grade=fig_grade,
    )
    st.download_button(
        "📄 Download Full Report (.docx)",
        docx_bytes,
        file_name="Result_Analysis_Report.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary",
        use_container_width=True,
    )
    st.download_button(
        "⬇️ Download Student Table Only (CSV)",
        student_df.to_csv(index=False).encode("utf-8"),
        file_name="student_result_table.csv",
        mime="text/csv",
    )

    st.markdown("---")
    st.caption(f"© {COPYRIGHT_NOTICE}")

elif not uploaded_pdf:
    st.info("⬅️ Fill in the examination details and upload the mark-list PDF from the sidebar, then click **Generate Report**.")
